#!/usr/bin/env python3
"""Create dataset-aware MADBase workbook templates from the frozen MNIST schemas."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from core import constants as C  # noqa: E402
from core.workbook import HEADER_ROW, validate_schema_map  # noqa: E402


TARGETS = (
    (C.MASTER_SPREADSHEET_PATH, C.MADBASE_MASTER_SPREADSHEET_PATH),
    (C.ABLATION_SPREADSHEET_PATH, C.MADBASE_ABLATION_SPREADSHEET_PATH),
    (C.MORE_REFERENCE_SPREADSHEET_PATH, C.MADBASE_MORE_REFERENCE_SPREADSHEET_PATH),
)
HEADER_FILL = PatternFill("solid", fgColor="FFD9EAF7")


def headers(ws):
    return [c.value for c in ws[HEADER_ROW] if c.value is not None]


def add_column(ws, name, value=None):
    hdr = headers(ws)
    if name in hdr:
        col = hdr.index(name) + 1
    else:
        col = len(hdr) + 1
        ws.cell(HEADER_ROW, col, name)
        ws.cell(HEADER_ROW, col).font = Font(bold=True)
        ws.cell(HEADER_ROW, col).fill = HEADER_FILL
    if value is not None:
        for row in range(HEADER_ROW + 1, ws.max_row + 1):
            if ws.cell(row, 1).value is not None:
                ws.cell(row, col, value)
    return col


def upsert_schema(wb, sheet, column, granularity, source, field, status="OK", note=""):
    ws = wb["00_Schema_Map"]
    for row in range(HEADER_ROW + 1, ws.max_row + 1):
        if ws.cell(row, 1).value == sheet and ws.cell(row, 2).value == column:
            vals = (granularity, source, field, status, note)
            for col, value in enumerate(vals, 3):
                ws.cell(row, col, value)
            return
    ws.append([sheet, column, granularity, source, field, status, note])


def prefix_value(value):
    if value is None:
        return value
    text = str(value)
    return text if text.startswith("AD-") else "AD-" + text


def adapt_prefilled_rows(wb):
    for sheet in ("01_Run_Summary", "02_Run_Config"):
        ws = wb[sheet]
        hdr = headers(ws)
        for row in range(HEADER_ROW + 1, ws.max_row + 1):
            if ws.cell(row, 1).value is None:
                continue
            values = {name: ws.cell(row, i + 1) for i, name in enumerate(hdr)}
            for uid_name in ("run_uid", "ablation_run_uid", "reference_run_uid",
                             "matching_run_uid_A", "matching_run_uid_D"):
                if uid_name in values and values[uid_name].value:
                    values[uid_name].value = prefix_value(values[uid_name].value)
            if "run_name" in values and values["run_name"].value:
                name = str(values["run_name"].value)
                if not name.startswith("arabic_digits/"):
                    values["run_name"].value = "arabic_digits/" + name
            if "split_manifest" in values and "seed" in values and values["seed"].value is not None:
                values["split_manifest"].value = (
                    f"research_arabic_digits/splits/confirmatory_seed{int(values['seed'].value)}.json"
                )


def build(source: Path, target: Path):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(source, target)
    wb = openpyxl.load_workbook(target)
    adapt_prefilled_rows(wb)
    for sheet in ("01_Run_Summary", "02_Run_Config"):
        ws = wb[sheet]
        add_column(ws, "dataset", C.DATASET_MADBASE)
        add_column(ws, "dataset_id", C.DATASET_IDS[C.DATASET_MADBASE])
        add_column(ws, "dataset_snapshot_sha256")
        upsert_schema(wb, sheet, "dataset", "per_run", "config.json (manifest run)", "dataset")
        upsert_schema(wb, sheet, "dataset_id", "per_run", "config.json (manifest run)", "dataset_id")
        upsert_schema(wb, sheet, "dataset_snapshot_sha256", "per_run", "config.json (manifest run)",
                      "dataset_snapshot_sha256")
    cfg = wb["02_Run_Config"]
    for name, field in (("dataset_license_policy", "dataset_license_policy"),
                        ("test_origin", "test_origin"), ("test_index_space", "test_index_space")):
        add_column(cfg, name)
        upsert_schema(wb, "02_Run_Config", name, "per_run", "config.json (manifest run)", field)

    pred = wb["11_Test_Predictions"]
    add_column(pred, "source_sample_id")
    add_column(pred, "source_index")
    upsert_schema(wb, "11_Test_Predictions", "sample_id", "per_test_sample",
                  "logs/test_sample_id.npy", "test_sample_id[i]", "OK", "stable source sample ID")
    upsert_schema(wb, "11_Test_Predictions", "source_sample_id", "per_test_sample",
                  "logs/test_sample_id.npy", "test_sample_id[i]", "OK", "stable source sample ID")
    upsert_schema(wb, "11_Test_Predictions", "source_index", "per_test_sample",
                  "logs/test_source_index.npy", "test_source_index[i]", "OK", "index in declared source pool")
    upsert_schema(wb, "11_Test_Predictions", "mnist_index", "per_test_sample", "N/A", "N/A", "N/A",
                  "legacy column retained blank in MADBase workbooks")

    rd = wb["00_README"]
    rd.append(["Dataset", "MADBASE-AUC-V1; Arabic-Indic handwritten digits; pinned archive checksums",
               "Dataset identity", "LOCKED"])
    rd.append(["Data policy", C.MADBASE_LICENSE_POLICY,
               "Official page has no explicit redistribution grant; raw files remain local", "LOCKED"])
    rd.append(["Split policy", "writer-disjoint train/validation; pilot holdout from train; fixed official-test subset",
               "Leakage prevention and paired evaluation", "LOCKED"])
    problems = validate_schema_map(wb)
    if problems:
        raise SystemExit("MADBase workbook schema invalid: " + "; ".join(problems[:20]))
    wb.save(target)
    print("ok", target.relative_to(ROOT))


def main():
    for source, target in TARGETS:
        build(ROOT / source, ROOT / target)


if __name__ == "__main__":
    main()
