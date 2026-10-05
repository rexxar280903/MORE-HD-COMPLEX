#!/usr/bin/env python3
"""Update the primary master workbook and (re)generate the ablation and MORE-reference
templates (G3-02, G3-06, G3-07). Idempotent: running it twice gives the same files.

* research_data/MORE_HD_master_confirmatory_240runs.xlsx  (primary A/D, updated in place)
* research_data/MORE_HD_master_ablation_90runs.xlsx       (B/C, ABL001-ABL018 x 5 seeds)
* research_data/MORE_HD_master_more_reference_120runs.xlsx (MORE-REPRO, MREF001-MREF024 x 5 seeds)
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

from core import constants as C  # noqa: E402
from core.config import ablation_condition_id, auto_generate_run_name, more_reference_condition_id, primary_condition_id  # noqa: E402
from core.seeds import manifest_seeds  # noqa: E402
from core.workbook import HEADER_ROW, validate_schema_map  # noqa: E402

PRIMARY = ROOT / C.MASTER_SPREADSHEET_PATH
ABLATION = ROOT / C.ABLATION_SPREADSHEET_PATH
MOREREF = ROOT / C.MORE_REFERENCE_SPREADSHEET_PATH
HEADER_FILL = PatternFill("solid", fgColor="FFD9EAF7")
CFG = "config.json (manifest run)"

# ---------------------------------------------------------------------------
# primary workbook updates
# ---------------------------------------------------------------------------
NEW_PRIMARY_COLUMNS = {
    "01_Run_Summary": [
        ("structural_diagnostics_sec", "per_run", "logs/stage_timing.jsonl", "wall_sec WHERE stage == 'structural_diagnostics'", "OK", "§9.4; tahap G2-05"),
        ("git_commit", "per_run", CFG, "environment.git_commit", "OK", "G4-01; commit kode yang menjalankan run"),
        ("backend_max_abs_diff", "per_run", CFG, "backend_self_check.max_abs_diff", "OK", "G4-01; selisih maksimum engine vs PennyLane default.qubit (harus <= 1e-10)"),
    ],
    "02_Run_Config": [
        ("correlation_rule", "per_run", CFG, "correlation_rule", "OK", "G4-04 (2026-10-04): MORE calc_class_rela, MSE/max, diagonal -1"),
        ("pca_svd_solver", "per_run", CFG, "pca_svd_solver", "OK", "2026-10-04: 'full' (SVD eksak); NULL bila bukan PCA"),
    ],
    "18_Circuit_Benchmark": [
        ("batch_size", "per_block_model", "benchmarks/circuit_microbenchmark_block<k>.json", "rows[].batch_size", "OK", "2026-10-04: satu repeat = satu forward pass batch ini (termasuk satu konstruksi U(theta)); waktu wall/cpu = per sampel"),
        ("unitary_median_sec", "per_block_model", "benchmarks/circuit_microbenchmark_block<k>.json", "rows[].unitary_median_sec", "OK", "Median waktu membangun kolom U(theta) sekali per objective evaluation"),
    ],
}
SCHEMA_FIXES = {
    ("01_Run_Summary", "status"): ("logs/run_status.json", "status", "OK", "G3-04: RUNNING/COMPLETED/FAILED; konsolidator hanya memakai COMPLETED"),
    ("01_Run_Summary", "attempt"): (CFG, "attempt", "OK", "G3-04: attempt baru hanya bila semua attempt sebelumnya FAILED"),
    ("02_Run_Config", "device_name"): (CFG, "device_name", "OK", "G4-01 dikunci 2026-10-04: core.engine.BatchedStatevector (diverifikasi vs PennyLane default.qubit)"),
    ("05_Correlation_Matrix", "raw_mse"): ("artifacts/correlation_raw_mse.npy", "raw_mse[i][j]", "OK", "G4-04: MSE rata-rata kanal antar mean kelas (100 sampel train/kelas)"),
    ("11_Test_Predictions", "sample_id"): ("artifacts/y_test.npy", "posisi baris (0-based)", "DERIVED", "G4-02: posisi dalam X_test_scaled.npy"),
    ("11_Test_Predictions", "mnist_index"): ("logs/test_mnist_index.npy", "test_mnist_index[i]", "OK", "G4-02: indeks MNIST test asli (urutan torchvision)"),
    ("11_Test_Predictions", "pred_class"): ("logs/test_predictions.npy", "y_pred[i]", "OK", "G3-02: label kuantum terdekat"),
    ("11_Test_Predictions", "distance_true"): ("logs/test_label_distances.npy", "D[i, true]", "OK", "G3-02: jarak cosine ke label kelas benar"),
    ("11_Test_Predictions", "distance_pred"): ("logs/test_label_distances.npy", "D[i, pred]", "OK", "G3-02: jarak cosine ke label terpilih"),
    ("11_Test_Predictions", "margin"): ("logs/test_margins.npy", "min_{k != true} D[i,k] - D[i,true]", "OK", "G3-02: positif = benar dengan margin"),
    ("11_Test_Predictions", "correct"): ("logs/test_predictions.npy", "true_class == pred_class", "DERIVED", ""),
}
README_ROWS = [
    ("Correlation matrix S (G4-04)", "MORE calc_class_rela: mean 100 sampel train/kelas, MSE rata-rata kanal, S_ij = MSE_ij / max MSE (i != j), diagonal -1", "Komparabilitas dengan MORE (keputusan 2026-10-04)", "LOCKED"),
    ("Simulation backend (G4-01)", "core.engine.BatchedStatevector: statevector analitik complex128, shots=None, 1 thread/proses; diverifikasi vs PennyLane default.qubit (<=1e-10) di unit test dan self-check tiap run", "Validitas structural zeros + kelayakan komputasi", "LOCKED"),
    ("MORE reproduction track", "MORE-REPRO (QCNN 8 qubit, 91 parameter, readout X/Y/Z, tanpa R) pada protokol identik; workbook research_data/MORE_HD_master_more_reference_120runs.xlsx", "Perbandingan langsung dengan MORE (keputusan 2026-10-04)", "LOCKED"),
    ("PCA solver", "svd_solver='full' (SVD eksak, fit train saja)", "Reproducibility (keputusan 2026-10-04)", "LOCKED"),
]


def style_header(ws, hdr_row=HEADER_ROW):
    for c in ws[hdr_row]:
        if c.value is not None:
            c.font = Font(bold=True)
            c.fill = HEADER_FILL
    ws.freeze_panes = f"A{hdr_row + 1}"


def headers(ws):
    return [c.value for c in ws[HEADER_ROW] if c.value is not None]


def add_column(ws, name, fill=None):
    hdr = headers(ws)
    if name in hdr:
        return hdr.index(name) + 1
    col = len(hdr) + 1
    cell = ws.cell(HEADER_ROW, col, name)
    cell.font = Font(bold=True)
    cell.fill = HEADER_FILL
    if fill is not None:
        for r in range(HEADER_ROW + 1, ws.max_row + 1):
            if ws.cell(r, 1).value is not None:
                ws.cell(r, col, fill(ws, r))
    return col


def schema_rows(wb):
    sm = wb["00_Schema_Map"]
    rows = {}
    for r in range(HEADER_ROW + 1, sm.max_row + 1):
        vals = [c.value for c in sm[r]][:7]
        if vals[0] is not None:
            rows[(vals[0], vals[1])] = (r, vals)
    return sm, rows


def update_primary():
    wb = openpyxl.load_workbook(PRIMARY)
    sm, rows = schema_rows(wb)
    for sheet, cols in NEW_PRIMARY_COLUMNS.items():
        ws = wb[sheet]
        for name, gran, src, field, status, note in cols:
            fill = None
            if sheet == "02_Run_Config" and name == "correlation_rule":
                fill = lambda ws_, r: C.CORRELATION_RULE  # noqa: E731
            if sheet == "02_Run_Config" and name == "pca_svd_solver":
                def fill(ws_, r):
                    fm = ws_.cell(r, headers(ws_).index("feature_method") + 1).value
                    return C.PCA_SVD_SOLVER if fm == "PCA" else None
            add_column(ws, name, fill)
            if (sheet, name) not in rows:
                sm.append([sheet, name, gran, src, field, status, note])
    ws = wb["02_Run_Config"]
    dcol = headers(ws).index("device_name") + 1
    for r in range(HEADER_ROW + 1, ws.max_row + 1):
        if ws.cell(r, 1).value is not None:
            ws.cell(r, dcol, C.DEVICE_NAME)
    sm, rows = schema_rows(wb)
    for key, (src, field, status, note) in SCHEMA_FIXES.items():
        r, _ = rows[key]
        sm.cell(r, 4, src)
        sm.cell(r, 5, field)
        sm.cell(r, 6, status)
        sm.cell(r, 7, note)
    rd = wb["00_README"]
    existing = {rd.cell(r, 1).value for r in range(1, rd.max_row + 1)}
    for row in README_ROWS:
        if row[0] not in existing:
            rd.append(list(row))
    for r in range(1, rd.max_row + 1):
        if rd.cell(r, 1).value == "Schema map":
            rd.cell(r, 4, "LOCKED")
    problems = validate_schema_map(wb)
    if problems:
        raise SystemExit("primary schema invalid: " + "; ".join(problems))
    wb.save(PRIMARY)
    return wb


# ---------------------------------------------------------------------------
# derived templates
# ---------------------------------------------------------------------------
TRACK_IDS = {"ABLATION": ("ablation_condition_id", "ablation_run_uid"),
             "MORE_REFERENCE": ("reference_condition_id", "reference_run_uid")}
EXTRA_01 = {
    "ABLATION": [
        ("model_code", "B = MORE-HD-60P, C = MORE-HD-C-FixedRZ"),
        ("matching_run_uid_A", "run primer MORE-HD seed/K/fitur sama (reuse, tidak dirun ulang)"),
        ("matching_run_uid_D", "run primer MORE-HD-C seed/K/fitur sama (reuse)"),
        ("n_trainable_params", "B 60, C 30"), ("n_fixed_rz_params", "B 0, C 30"),
        ("n_variational_layers", "B 6, C 3"), ("variational_gate_count", "jumlah RY+RZ di ansatz"),
        ("entangling_block_count", "B 6, C 3"), ("cnot_count", "B 66, C 33"), ("circuit_depth", "kedalaman sirkuit (encoding termasuk)"),
        ("ry_core_seed", "init:ry_core (dibagi A/B/C/D)"), ("ry_extra_seed", "B saja"), ("rz_phase_seed", "C saja (dibagi dengan D)"),
        ("fixed_rz_sha256", "C saja; hash vektor RZ beku"), ("paired_pair_manifest_sha256", "harus sama dengan A/D cell yang sama"),
        ("paired_split_manifest_sha256", "harus sama dengan A/D seed yang sama"),
        ("paired_input_check", "MATCH bila array input identik dengan run A; A_NOT_AVAILABLE bila A belum ada"),
    ],
    "MORE_REFERENCE": [
        ("model_code", "M = MORE-REPRO"), ("matching_run_uid_A", "run primer MORE-HD seed/K/fitur sama"),
        ("matching_run_uid_D", "run primer MORE-HD-C seed/K/fitur sama"),
        ("n_trainable_params", "91 (QCNN MORE)"), ("n_fixed_rz_params", "0"), ("n_variational_layers", "3 conv+pool"),
        ("variational_gate_count", "jumlah RX/RY/RZ terparameter"), ("entangling_block_count", "0 (QCNN)"),
        ("cnot_count", "7 (pooling)"), ("circuit_depth", "kedalaman sirkuit (encoding termasuk)"),
        ("ry_core_seed", "tidak dipakai M; dicatat untuk keseragaman"), ("ry_extra_seed", "tidak dipakai"),
        ("rz_phase_seed", "tidak dipakai"), ("more_init_seed", "init:more_qcnn, U[0,1) seperti qiskit-machine-learning"),
        ("fixed_rz_sha256", "tidak dipakai"), ("paired_pair_manifest_sha256", "harus sama dengan A/D cell yang sama"),
        ("paired_split_manifest_sha256", "harus sama dengan A/D seed yang sama"),
        ("paired_input_check", "MATCH bila array input identik dengan run A"),
    ],
}
COMPARISON_SHEETS = {
    "ABLATION": ("20_Ablation_Contrasts",
                 "Planned contrasts A-B, A-C, B-D, C-D per (K, feature, seed); A/D dari workbook primer (SAP §6.1)",
                 ["K", "classes", "feature_method", "seed", "run_uid_A", "run_uid_B", "run_uid_C", "run_uid_D",
                  "pairing_check", "accuracy_A", "accuracy_B", "accuracy_C", "accuracy_D", "f1_A", "f1_B", "f1_C", "f1_D",
                  "min_sep_ratio_A", "min_sep_ratio_B", "min_sep_ratio_C", "min_sep_ratio_D",
                  "yodd_frac_A", "yodd_frac_B", "yodd_frac_C", "yodd_frac_D",
                  "delta_acc_AB", "delta_acc_AC", "delta_acc_BD", "delta_acc_CD",
                  "delta_f1_AB", "delta_f1_AC", "delta_f1_BD", "delta_f1_CD",
                  "delta_msr_AB", "delta_msr_AC", "delta_msr_BD", "delta_msr_CD"]),
    "MORE_REFERENCE": ("21_MORE_Comparison",
                       "MORE-REPRO (M) vs MORE-HD (A) vs MORE-HD-C (D) pada split, pair set, budget identik; plus Tabel I MORE (deskriptif)",
                       ["K", "classes", "feature_method", "seed", "run_uid_M", "run_uid_A", "run_uid_D", "pairing_check",
                        "accuracy_M", "accuracy_A", "accuracy_D", "f1_M", "f1_A", "f1_D",
                        "min_sep_M", "min_sep_A", "min_sep_D", "min_sep_ratio_M", "min_sep_ratio_A", "min_sep_ratio_D",
                        "delta_acc_A_minus_M", "delta_acc_D_minus_M", "delta_f1_A_minus_M", "delta_f1_D_minus_M",
                        "delta_msr_A_minus_M", "delta_msr_D_minus_M",
                        "MORE_TableI_without_R_pct", "MORE_TableI_with_R_pct"]),
}


def cells_for(track):
    out = []
    ks = C.ABLATION_K_VALUES if track == "ABLATION" else C.K_VALUES
    models = C.ABLATION_MODELS if track == "ABLATION" else (C.ARCH_MORE_REPRO,)
    for seed in C.CONFIRMATORY_SEEDS:
        for k in ks:
            for fm in C.FEATURE_METHODS:
                for model in models:
                    out.append((seed, k, fm, model))
    return out


def run_name_for(track, seed, k, fm, model):
    cid = ablation_condition_id(k, fm, model) if track == "ABLATION" else more_reference_condition_id(k, fm)
    sub = "ablation" if track == "ABLATION" else "more_reference"
    cls = "-".join(str(c) for c in range(k))
    return cid, f"{cid}-S{seed}", f"{sub}/{cid}-S{seed}_cls-{cls}_{fm}_{model}"


def build_derived(track, primary_wb, path):
    cid_col, uid_col = TRACK_IDS[track]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    p_sm, p_rows = schema_rows(primary_wb)
    schema = []

    def map_col(sheet, primary_sheet, col):
        src_col = {"condition_id": "condition_id", "run_uid": "run_uid"}
        if col in (cid_col, uid_col):
            base = "condition_id" if col == cid_col else "run_uid"
            note = "ID track terpisah dari R001-R048 (tidak menduplikasi A/D)"
            schema.append([sheet, col, "per_row", CFG, base, "TEMPLATE" if sheet in ("01_Run_Summary", "02_Run_Config") else "OK", note])
            return
        key = (primary_sheet, src_col.get(col, col))
        if key in p_rows:
            vals = list(p_rows[key][1])
            vals[0] = sheet
            schema.append(vals)
        else:
            schema.append([sheet, col, "per_run", CFG, col, "OK", ""])

    # README
    rd = wb.create_sheet("00_README")
    title = {"ABLATION": "MORE-HD-C Targeted Ablation Workbook (90 runs: B/C x K{3,6,10} x 3 fitur x 5 seed)",
             "MORE_REFERENCE": "MORE Reproduction Workbook (120 runs: MORE-REPRO x K3..10 x 3 fitur x 5 seed)"}[track]
    rd.append([title])
    rd["A1"].font = Font(bold=True, size=14)
    rd.append([])
    rd.append(["Protocol item", "Locked value / rule", "Purpose", "Status"])
    for c in rd[3]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
    common = [
        ("Seeds", "101, 202, 303, 404, 505 (seed 42 PILOT ONLY)", "Paired replications", "LOCKED"),
        ("Data", "1000/100/200 per kelas; split, pair set, S, budget dan evaluasi identik dengan run primer", "Paired comparison", "LOCKED"),
        ("Primary rows", "Run A/D tidak diduplikasi; ditautkan lewat matching_run_uid_A/D", "SAP §12", "LOCKED"),
    ]
    if track == "ABLATION":
        common.append(("Models", "B = MORE-HD-60P (6 layer RY, 60 trainable); C = MORE-HD-C-FixedRZ (30 RY trainable + 30 RZ beku)", "G1-05/G1-07", "LOCKED"))
    else:
        common.append(("Model", "M = MORE-REPRO: QCNN MORE 8 qubit, ZFeatureMap(reps=2) pada 2x fitur, 91 parameter, readout qubit 7 X/Y/Z, tanpa loss adjuster R", "Perbandingan langsung dengan MORE", "LOCKED"))
    for row in common:
        rd.append(list(row))

    # run sheets cloned from the primary
    for name in ("01_Run_Summary", "02_Run_Config", "03_Clustering_History", "04_Supervised_History",
                 "05_Correlation_Matrix", "06_Quantum_Labels", "07_Pairwise_Distance", "08_Dimension_Activity",
                 "09_Per_Class_Metrics", "10_Confusion_Data", "11_Test_Predictions", "19_Threshold_Sensitivity"):
        src = primary_wb[name]
        hdr = [h for h in headers(src)]
        hdr = [cid_col if h == "condition_id" else uid_col if h == "run_uid" else h for h in hdr]
        if name == "06_Quantum_Labels" and track == "MORE_REFERENCE":
            i0 = hdr.index("IX")
            hdr = hdr[:i0] + list(C.OBSERVABLE_LABELS_MORE) + hdr[i0 + 15:]
        if name == "01_Run_Summary":
            hdr = hdr + [c for c, _ in EXTRA_01[track] if c not in hdr]
        ws = wb.create_sheet(name)
        ws.append([f"{src.cell(1, 1).value} — {track}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        ws.append(hdr)
        style_header(ws)
        for col in hdr:
            if name == "06_Quantum_Labels" and col in C.OBSERVABLE_LABELS_MORE:
                schema.append([name, col, "per_class", "artifacts/quantum_labels.json", f"quantum_labels[class][{col}]", "OK", "Readout MORE: qubit 7, urutan X, Y, Z"])
                continue
            extra = dict(EXTRA_01[track])
            if name == "01_Run_Summary" and col in extra:
                field = {"circuit_depth": "circuit_resources.depth", "fixed_rz_sha256": "initialization_manifest.fixed_rz_sha256",
                         "paired_split_manifest_sha256": "split_manifest_sha256",
                         "paired_input_check": "paired_input_check.status"}.get(col, col)
                schema.append([name, col, "per_run", CFG, field, "OK", extra[col]])
                continue
            map_col(name, name, col)
        if name in ("01_Run_Summary", "02_Run_Config"):
            for seed, k, fm, model in cells_for(track):
                cid, uid, rname = run_name_for(track, seed, k, fm, model)
                vals = {cid_col: cid, uid_col: uid, "run_name": rname, "run_mode": "CONFIRMATORY", "K": k,
                        "classes": ",".join(str(c) for c in range(k)), "feature_method": fm, "architecture": model,
                        "seed": seed, "split_manifest": f"splits/seed{seed}.json"}
                if name == "01_Run_Summary":
                    n_par = {C.ARCH_MORE_HD_60P: 60, C.ARCH_MORE_HD_C_FIXED_RZ: 30, C.ARCH_MORE_REPRO: 91}[model]
                    vals.update({"status": "NOT_STARTED", "replication_id": f"S{seed}", "n_params": n_par,
                                 "n_train_total": 1000 * k, "n_val_total": 100 * k, "n_test_total": 200 * k, "attempt": 1,
                                 "model_code": C.MODEL_CODE[model],
                                 "matching_run_uid_A": f"{primary_condition_id(k, fm, C.ARCH_MORE_HD)}-S{seed}",
                                 "matching_run_uid_D": f"{primary_condition_id(k, fm, C.ARCH_MORE_HD_C)}-S{seed}"})
                else:
                    s = manifest_seeds(seed)
                    vals.update({"n_train_per_class": 1000, "n_val_per_class": 100, "n_test_per_class": 200,
                                 "n_data_qubits": 8, "n_readout_qubits": 1 if model == C.ARCH_MORE_REPRO else 2,
                                 "n_input_channels": 8, "cobyla_tol": C.COBYLA_TOL, "cobyla_rhobeg": C.COBYLA_RHOBEG,
                                 "master_seed": seed, "data_seed": s["data_seed"], "ry_core_seed": s["ry_core_seed"],
                                 "rz_phase_seed": s["rz_phase_seed"], "pair_seed": s["pair_seed"], "label_seed": s["label_seed"],
                                 "paired_split_key": f"seed{seed}_K{k}", "nested_class_split": True,
                                 "active_dim_threshold": C.ACTIVE_DIM_THRESHOLD, "subseed_rule": C.SUBSEED_RULE,
                                 "simulation_mode": C.SIMULATION_MODE, "sim_dtype": C.SIM_DTYPE, "device_name": C.DEVICE_NAME,
                                 "eps_norm": C.EPS_NORM, "centroid_rule": C.CENTROID_RULE, "centroid_source": C.CENTROID_SOURCE,
                                 "val_monitor_policy": C.VAL_MONITOR_POLICY, "correlation_rule": C.CORRELATION_RULE,
                                 "pca_svd_solver": C.PCA_SVD_SOLVER if fm == "PCA" else None,
                                 "n_params": {C.ARCH_MORE_HD_60P: 60, C.ARCH_MORE_HD_C_FIXED_RZ: 30, C.ARCH_MORE_REPRO: 91}[model],
                                 "pca_n_components": 8 if fm == "PCA" else None, "hu_raw_dim": 7 if fm == "HU" else None,
                                 "hu_padding_value": 0.0 if fm == "HU" else None,
                                 "zernike_terms": str([list(t) for t in C.ZERNIKE_TERMS]) if fm == "ZERNIKE" else None})
                ws.append([vals.get(h) for h in hdr])

    # comparison sheet
    sheet, title, hdr = COMPARISON_SHEETS[track]
    ws = wb.create_sheet(sheet)
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(hdr)
    style_header(ws)
    seen = set()
    for seed, k, fm, _ in cells_for(track):
        if (seed, k, fm) in seen:
            continue
        seen.add((seed, k, fm))
        row = {"K": k, "classes": ",".join(str(c) for c in range(k)), "feature_method": fm, "seed": seed,
               "run_uid_A": f"{primary_condition_id(k, fm, C.ARCH_MORE_HD)}-S{seed}",
               "run_uid_D": f"{primary_condition_id(k, fm, C.ARCH_MORE_HD_C)}-S{seed}"}
        if track == "ABLATION":
            row["run_uid_B"] = f"{ablation_condition_id(k, fm, C.ARCH_MORE_HD_60P)}-S{seed}"
            row["run_uid_C"] = f"{ablation_condition_id(k, fm, C.ARCH_MORE_HD_C_FIXED_RZ)}-S{seed}"
        else:
            row["run_uid_M"] = f"{more_reference_condition_id(k, fm)}-S{seed}"
            row["MORE_TableI_without_R_pct"], row["MORE_TableI_with_R_pct"] = C.MORE_TABLE_I[k]
        ws.append([row.get(h) for h in hdr])
    for col in hdr:
        status = "TEMPLATE" if col in ("K", "classes", "feature_method", "seed") or col.startswith("run_uid") else "DERIVED"
        if col.startswith("MORE_TableI"):
            status = "STATIC"
        src = "analysis/run_analysis.py" if status == "DERIVED" else ("Wu et al. (2023) Tabel I" if status == "STATIC" else "template (prefilled)")
        schema.append([sheet, col, "per_cell_seed", src, col, status, "Diisi analisis; deskriptif + uji SAP" if status == "DERIVED" else ""])

    sm = wb.create_sheet("00_Schema_Map", 1)
    sm.append([f"Schema Map — {track} (sumber kebenaran kolom → artefak → field; G3-02/G3-07)"])
    sm["A1"].font = Font(bold=True, size=14)
    sm.append([])
    sm.append(["sheet", "column", "granularity", "source_artifact", "source_field", "status", "note"])
    style_header(sm)
    for row in schema:
        sm.append(row)
    problems = validate_schema_map(wb)
    if problems:
        raise SystemExit(f"{track} schema invalid: " + "; ".join(problems[:20]))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main():
    primary = update_primary()
    build_derived("ABLATION", primary, ABLATION)
    build_derived("MORE_REFERENCE", primary, MOREREF)
    for p in (PRIMARY, ABLATION, MOREREF):
        print("ok", p.relative_to(ROOT))


if __name__ == "__main__":
    main()
