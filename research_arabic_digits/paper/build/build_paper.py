"""Render naskah MADBase (bahasa Indonesia) ke template Word VJCS.

Salinan terpisah dari paper/build/build_paper.py studi MNIST (file MNIST tidak disentuh). Perbedaan:
template dibaca dari research_arabic_digits/paper/template/, keluaran ke research_arabic_digits/paper/, dan
persamaan LaTeX dikonversi ke OMML memakai MML2OMML.XSL milik Microsoft Office bila tersedia (Windows),
atau pandoc bila tidak (Linux/macOS).

Usage:  python research_arabic_digits/paper/build/build_paper.py   (menulis MORE-HD-C_MADBase_paper_ID.docx)

Content lives in content_en.py / content_id.py as a list of blocks; this file only renders them:
  ("h1"|"h2"|"h3", text)        numbered headings (template styles Heading 1-3)
  ("h1n", text)                 unnumbered heading (Acknowledgments, References)
  ("app", text) / ("app2", text) appendix headings (template styles Appendix 1/2)
  ("p", markup)                 paragraph; first after a heading is flush left (Text), later ones indented
  ("bul", [markup, ...])        bullet list (template style Bullet List)
  ("eq", label, latex)          numbered display equation, number flush right
  ("fig", label, file, width_in, caption_markup)
  ("tab", label, caption_markup, header, rows, col_widths_twips)
  ("refs",)                     reference list in order of first citation
Inline markup: **bold**, *italic*, $latex$, [@key,@key] superscript citations, [#key] "Ref. n",
{eq:x} {fig:x} {tab:x} cross references.
"""
import copy
import importlib
import json
import re
import sys
from pathlib import Path

import latex2mathml.converter as l2m
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
PAPER = HERE.parent
TEMPLATE = PAPER / "template" / "ws-vjcs template.docx"
OUT_NAME = "MORE-HD-C_MADBase_paper_{LANG}.docx"
_XSL_PATH = Path(r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL")
XSL = etree.XSLT(etree.parse(str(_XSL_PATH))) if _XSL_PATH.exists() else None
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

sys.path.insert(0, str(HERE))


# ------------------------------------------------------------------ math
FUNC_NAMES = {"min", "max", "med", "sign", "arg", "log", "cos", "Metric", "MSE", "exp"}


_PANDOC_CACHE = {}


def _omml_pandoc(latex: str):
    """LaTeX -> OMML via pandoc (texmath); hasil di-cache per string."""
    import subprocess
    import tempfile
    import zipfile
    if latex not in _PANDOC_CACHE:
        with tempfile.TemporaryDirectory() as tmp:
            md, dx = Path(tmp) / "m.md", Path(tmp) / "m.docx"
            md.write_text("$" + latex + "$\n", encoding="utf-8")
            subprocess.run(["pandoc", str(md), "-o", str(dx)], check=True)
            root = etree.fromstring(zipfile.ZipFile(dx).read("word/document.xml"))
            node = next(root.iter("{%s}oMath" % M_NS))
            _PANDOC_CACHE[latex] = etree.tostring(node)
    return etree.fromstring(_PANDOC_CACHE[latex])


def omml(latex: str, display: bool = False):
    M = "{%s}" % M_NS
    if XSL is not None:
        mathml = l2m.convert(latex)
        out = XSL(etree.fromstring(mathml)).getroot()
        # Word wants the math elements without the leftover mml namespace declaration
        out = etree.fromstring(etree.tostring(out).replace(b' xmlns:mml="http://www.w3.org/1998/Math/MathML"', b""))
    else:
        out = _omml_pandoc(latex)
    # upright function names (min, max, med, ...)
    for r in out.iter(M + "r"):
        t = r.find(M + "t")
        if t is not None and t.text and t.text.strip() in FUNC_NAMES and r.find(M + "rPr") is None:
            rpr = etree.Element(M + "rPr")
            sty = etree.SubElement(rpr, M + "sty")
            sty.set(M + "val", "p")
            r.insert(0, rpr)
    # n-ary operators converted with an empty body: move the following siblings into the body
    for nary in list(out.iter(M + "nary")):
        e = nary.find(M + "e")
        if e is not None and len(e) == 0:
            parent = nary.getparent()
            idx = list(parent).index(nary)
            for sib in list(parent)[idx + 1:]:
                e.append(sib)
        if display:
            pr = nary.find(M + "naryPr")
            if pr is None:
                pr = etree.Element(M + "naryPr")
                nary.insert(0, pr)
            for old in pr.findall(M + "limLoc"):
                pr.remove(old)
            ll = etree.Element(M + "limLoc")
            ll.set(M + "val", "undOvr")
            pr.insert(1 if pr.find(M + "chr") is not None else 0, ll)
    return out


# ------------------------------------------------------------------ references
def author_list(authors, max_all=5):
    def initials(given):
        parts = []
        for chunk in given.replace(".", ". ").split():
            sub = [s for s in chunk.split("-") if s]
            parts.append("-".join(s[0] + "." for s in sub))
        return " ".join(parts)

    names = [f"{initials(g)} {f}".strip() for g, f in authors]
    if len(names) > max_all:
        return names[0] + " et al."
    if len(names) == 1:
        return names[0]
    return ", ".join(names[:-1]) + " and " + names[-1]


def load_refs(lang):
    cr = json.load(open(HERE / "crossref_refs.json"))
    from references import REFS
    out = {}
    for key, spec in REFS.items():
        if "doi" in spec and spec["doi"] in cr and "text" not in spec:
            c = cr[spec["doi"]]
            au = spec.get("authors") or author_list(c["authors"])
            jr = spec.get("journal", c["short"] or c["journal"])
            pages = c["page"] or ""
            pages = pages.replace("-", "\u2013")
            title = spec.get("title", c["title"]).rstrip("*")
            out[key] = f"{au}, {title}, {jr} {c['vol']} ({c['year']}) {pages}, doi:{spec['doi']}."
        else:
            out[key] = spec["text"]
    return out


# ------------------------------------------------------------------ docx helpers
def clear_paragraph(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def set_header_text(doc, mapping):
    for rel in doc.part.rels.values():
        if "header" in rel.reltype:
            part = rel.target_part
            root = part._element
            for t in root.iter(qn("w:t")):
                for old, new in mapping.items():
                    if t.text and old in t.text:
                        t.text = t.text.replace(old, new)
            # collapse split runs: replace whole-text matches spanning runs
            texts = list(root.iter(qn("w:t")))
            joined = "".join(t.text or "" for t in texts)
            for old, new in mapping.items():
                if old in joined and texts:
                    texts[0].text = joined.replace(old, new)
                    for t in texts[1:]:
                        t.text = ""
                    joined = texts[0].text


class Renderer:
    def __init__(self, content, lang):
        self.c = content
        self.lang = lang
        self.L = content.LABELS
        self.refs_text = load_refs(lang)
        self.cite_order = []
        self.numbers = {}
        self._number_items()

    # first pass: number equations, figures, tables, appendix equations; collect citation order
    def _number_items(self):
        counters = {"eq": 0, "fig": 0, "tab": 0}
        appendix = None
        app_count = {"eq": 0, "fig": 0, "tab": 0}
        letters = "ABCDEFGH"
        n_app = -1
        for b in self.c.BLOCKS:
            kind = b[0]
            if kind == "app":
                n_app += 1
                appendix = letters[n_app]
                app_count = {"eq": 0, "fig": 0, "tab": 0}
            if kind in ("eq", "fig", "tab"):
                if appendix:
                    app_count[kind] += 1
                    self.numbers[f"{kind}:{b[1]}"] = f"{appendix}.{app_count[kind]}"
                else:
                    counters[kind] += 1
                    self.numbers[f"{kind}:{b[1]}"] = str(counters[kind])
        # citation order by first appearance in text (abstract excluded per template)
        texts = []
        for b in self.c.BLOCKS:
            if b[0] in ("p", "pn"):
                texts.append(b[1])
            elif b[0] == "bul":
                texts.extend(b[1])
            elif b[0] == "fig":
                texts.append(b[4])
            elif b[0] == "tab":
                texts.append(b[2])
        for t in texts:
            for grp in re.findall(r"\[[@#]([^\]]+)\]", t):
                for k in re.split(r"\s*,\s*@?", grp):
                    k = k.lstrip("@#")
                    if k not in self.cite_order:
                        if k not in self.refs_text:
                            raise KeyError(f"unknown reference key {k}")
                        self.cite_order.append(k)

    def cite_numbers(self, keys):
        nums = sorted(self.cite_order.index(k) + 1 for k in keys)
        # compress ranges 1,2,3 -> 1-3
        out, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            out.append(f"{nums[i]}\u2013{nums[j]}" if j - i >= 2 else ",".join(str(n) for n in nums[i:j + 1]))
            i = j + 1
        return ",".join(out)

    def resolve_refs(self, text):
        def rep(m):
            key = m.group(1)
            if key not in self.numbers:
                raise KeyError(f"unknown cross reference {key}")
            return self.numbers[key]
        return re.sub(r"\{((?:eq|fig|tab):[A-Za-z0-9_]+)\}", rep, text)

    # inline markup -> runs / math
    def add_inline(self, p, text, base_italic=False):
        text = self.resolve_refs(text)
        tokens = re.split(r"(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*]+\*|\[@[^\]]+\]|\[#[^\]]+\])", text)
        for tok in tokens:
            if not tok:
                continue
            if tok.startswith("$"):
                p._p.append(omml(tok[1:-1]))
            elif tok.startswith("**"):
                r = p.add_run(tok[2:-2])
                r.bold = True
            elif tok.startswith("*"):
                r = p.add_run(tok[1:-1])
                r.italic = not base_italic
            elif tok.startswith("[@"):
                keys = [k.strip().lstrip("@") for k in tok[2:-1].split(",")]
                r = p.add_run(self.cite_numbers(keys))
                r.font.superscript = True
            elif tok.startswith("[#"):
                keys = [k.strip().lstrip("#") for k in tok[2:-1].split(",")]
                word = self.L["ref"] if len(keys) == 1 else self.L["refs"]
                p.add_run(f"{word}\u00a0{self.cite_numbers(keys)}")
            else:
                r = p.add_run(tok)
                if base_italic:
                    r.italic = True
        return p

    # ---------------------------------------------------------- blocks
    def heading(self, doc, style, text, numbered=True):
        p = doc.add_paragraph(style=style)
        if not numbered:
            pPr = p._p.get_or_add_pPr()
            numPr = OxmlElement("w:numPr")
            ilvl = OxmlElement("w:ilvl")
            ilvl.set(qn("w:val"), "0")
            numId = OxmlElement("w:numId")
            numId.set(qn("w:val"), "0")
            numPr.append(ilvl)
            numPr.append(numId)
            pPr.append(numPr)
        self.add_inline(p, text)
        return p

    def equation(self, doc, label, latex):
        """Display equation centred on the text width with its number flush right (borderless 3-cell table)."""
        widths = [720, 5760, 720]
        t = doc.add_table(rows=1, cols=3)
        tblPr = t._tbl.tblPr
        for tag, attrs in (("w:tblW", {"w:w": "7200", "w:type": "dxa"}), ("w:jc", {"w:val": "center"}),
                           ("w:tblLayout", {"w:type": "fixed"})):
            old = tblPr.find(qn(tag))
            if old is not None:
                tblPr.remove(old)
            e = OxmlElement(tag)
            for k, v in attrs.items():
                e.set(qn(k), v)
            tblPr.append(e)
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "nil")
            borders.append(b)
        tblPr.append(borders)
        mar = OxmlElement("w:tblCellMar")
        for side in ("left", "right"):
            m = OxmlElement(f"w:{side}")
            m.set(qn("w:w"), "0")
            m.set(qn("w:type"), "dxa")
            mar.append(m)
        tblPr.append(mar)
        for gc, w in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        row = t.rows[0]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        for cell, w in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = OxmlElement("w:tcW")
            tcW.set(qn("w:w"), str(w))
            tcW.set(qn("w:type"), "dxa")
            tcPr.append(tcW)
            va = OxmlElement("w:vAlign")
            va.set(qn("w:val"), "center")
            tcPr.append(va)
            cell.paragraphs[0].style = doc.styles["Equation"]
        mid = row.cells[1].paragraphs[0]
        M = "{%s}" % M_NS
        para = etree.SubElement(mid._p, M + "oMathPara")
        ppr = etree.SubElement(para, M + "oMathParaPr")
        jc = etree.SubElement(ppr, M + "jc")
        jc.set(M + "val", "center")
        para.append(omml(latex, display=True))
        num = row.cells[2].paragraphs[0]
        num.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        num.add_run(f"({self.numbers['eq:' + label]})")

    def figure(self, doc, label, fname, width, caption):
        p = doc.add_paragraph(style="Figure")
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
        p.add_run().add_picture(str(PAPER / "figures" / self.lang / fname), width=Inches(width))
        cap = doc.add_paragraph(style="Figure Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        cap.add_run(f"{self.L['fig']}\u00a0{self.numbers['fig:' + label]}.\u2002")
        self.add_inline(cap, caption)

    def table(self, doc, label, caption, header, rows, widths):
        cap = doc.add_paragraph(style="Table Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if len(caption) > 90 else WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run(f"{self.L['tab']}\u00a0{self.numbers['tab:' + label]}.\u2002")
        self.add_inline(cap, caption)
        n_head = len(header)
        all_rows = header + rows
        ncol = len(widths)
        t = doc.add_table(rows=len(all_rows), cols=ncol)
        tbl = t._tbl
        tblPr = tbl.tblPr
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        tblPr.append(jc)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        tblW.set(qn("w:w"), str(sum(widths)))
        tblW.set(qn("w:type"), "dxa")
        grid = tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
            gc.set(qn("w:w"), str(w))
        for ri, row in enumerate(all_rows):
            tr = t.rows[ri]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit")
            trPr.append(cs)
            if ri < n_head:
                th = OxmlElement("w:tblHeader")
                trPr.append(th)
            ci = 0
            cells = list(row)
            for cj in range(ncol):
                cell = tr.cells[cj]
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), str(widths[cj]))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)
                borders = OxmlElement("w:tcBorders")
                def border(side, sz):
                    e = OxmlElement(f"w:{side}")
                    e.set(qn("w:val"), "single")
                    e.set(qn("w:sz"), str(sz))
                    e.set(qn("w:space"), "0")
                    e.set(qn("w:color"), "auto")
                    borders.append(e)
                if ri == 0:
                    border("top", 8)
                if ri == n_head - 1:
                    border("bottom", 4)
                if ri == len(all_rows) - 1:
                    border("bottom", 8)
                tcPr.append(borders)
                va = OxmlElement("w:vAlign")
                va.set(qn("w:val"), "center")
                tcPr.append(va)
                txt = cells[cj] if cj < len(cells) else ""
                para = cell.paragraphs[0]
                para.style = doc.styles["Table"]
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT if cj == 0 else WD_ALIGN_PARAGRAPH.CENTER
                lines = str(txt).split("\n")
                for li, line in enumerate(lines):
                    if li:
                        para = cell.add_paragraph(style="Table")
                        para.alignment = WD_ALIGN_PARAGRAPH.LEFT if cj == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    self.add_inline(para, line)
                    if ri < n_head:
                        for run in para.runs:
                            run.bold = True
        # small spacer after table
        doc.add_paragraph(style="Text")

    def references(self, doc):
        for k in self.cite_order:
            p = doc.add_paragraph(style="Reference")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.add_run(self.refs_text[k])

    # ---------------------------------------------------------- front matter
    def front_matter(self, doc):
        ps = doc.paragraphs
        title, auth1, aff1, auth2, aff2, hist, abst, kw1, kw2 = ps[:9]
        clear_paragraph(title)
        for i, line in enumerate(self.c.TITLE_LINES):
            r = title.add_run(line)
            if i < len(self.c.TITLE_LINES) - 1:
                r.add_break()
        for p, txt in ((auth1, self.c.AUTHORS[0]), (auth2, self.c.AUTHORS[1])):
            clear_paragraph(p)
            p.add_run(txt)
        for p, lines in ((aff1, self.c.AFFILIATIONS[0]), (aff2, self.c.AFFILIATIONS[1])):
            clear_paragraph(p)
            for i, line in enumerate(lines):
                r = p.add_run(line)
                if i < len(lines) - 1:
                    r.add_break()
        clear_paragraph(hist)
        for i, line in enumerate(self.c.HISTORY):
            r = hist.add_run(line)
            if i < len(self.c.HISTORY) - 1:
                r.add_break()
        clear_paragraph(abst)
        self.add_inline(abst, self.c.ABSTRACT)
        # kw1 removed; kw2 keeps the section break of the front matter
        kw1._p.getparent().remove(kw1._p)
        clear_paragraph(kw2)
        self.add_inline(kw2, self.c.KEYWORDS)

    def build(self, out_path):
        doc = Document(str(TEMPLATE))
        body = doc.element.body
        children = list(body)
        # keep first 9 paragraphs (front matter incl. section break) and the final sectPr
        kept = 0
        for el in children:
            if el.tag == qn("w:p") and kept < 9:
                kept += 1
                continue
            if el.tag == qn("w:sectPr"):
                continue
            body.remove(el)
        self.front_matter(doc)
        set_header_text(doc, self.c.HEADERS)
        if self.lang == "id":
            self.patch_appendix_label(doc)

        prev = "heading"
        for b in self.c.BLOCKS:
            kind = b[0]
            if kind in ("h1", "h2", "h3"):
                self.heading(doc, {"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[kind], b[1])
                prev = "heading"
            elif kind == "h1n":
                self.heading(doc, "Heading 1", b[1], numbered=False)
                prev = "heading"
            elif kind == "app":
                self.heading(doc, "Appendix 1", b[1])
                prev = "heading"
            elif kind == "app2":
                self.heading(doc, "Appendix 2", b[1])
                prev = "heading"
            elif kind == "p":
                style = "Text" if prev in ("heading", "fig", "tab", "eq", "bul") and prev == "heading" else "Text Indent"
                if prev in ("eq",) and b[1].startswith(self.L["where"]):
                    style = "Text"
                p = doc.add_paragraph(style=style)
                self.add_inline(p, b[1])
                prev = "p"
            elif kind == "pn":  # paragraph never indented
                p = doc.add_paragraph(style="Text")
                self.add_inline(p, b[1])
                prev = "p"
            elif kind == "bul":
                for item in b[1]:
                    p = doc.add_paragraph(style="Bullet List")
                    self.add_inline(p, item)
                prev = "bul"
            elif kind == "eq":
                self.equation(doc, b[1], b[2])
                prev = "eq"
            elif kind == "fig":
                self.figure(doc, *b[1:])
                prev = "fig"
            elif kind == "tab":
                self.table(doc, *b[1:])
                prev = "tab"
            elif kind == "refs":
                self.references(doc)
            else:
                raise ValueError(kind)
        props = doc.core_properties
        props.title = " ".join(self.c.TITLE_LINES)
        props.author = self.c.AUTHORS[0]
        props.subject = "MORE-HD / MORE-HD-C pada MADBase (digit Arab)"
        doc.save(str(out_path))
        print("wrote", out_path, "| refs cited:", len(self.cite_order), "| numbers:", len(self.numbers))

    @staticmethod
    def patch_appendix_label(doc):
        numbering = doc.part.numbering_part.element
        for lt in numbering.iter(qn("w:lvlText")):
            v = lt.get(qn("w:val"))
            if v and v.startswith("Appendix"):
                lt.set(qn("w:val"), v.replace("Appendix", "Lampiran"))


if __name__ == "__main__":
    langs = sys.argv[1:] or ["id"]
    for lang in langs:
        content = importlib.import_module(f"content_{lang}")
        out = PAPER / OUT_NAME.format(LANG=lang.upper())
        Renderer(content, lang).build(out)
