"""Workbook rows from run artifacts, per-run workbooks, schema-map validation and
the idempotent consolidator (pseudocode §0.2, §1 UPDATE_LOCAL_RUN_SPREADSHEET;
G3-01, G3-02, G3-03, G3-06, G3-07).

Artifacts (JSON/JSONL/NPY/BIN) remain the primary data; every workbook value
is derived from them through the row builders below, so a per-run workbook and
the consolidated workbook contain identical values for a run.
"""

from __future__ import annotations

import shutil
from functools import cached_property
from pathlib import Path

import numpy as np
import openpyxl

from . import constants as C
from .artifact_io import load_json, load_npy, read_jsonl, sha256_file
from .numerics import pairwise_label_distances

HEADER_ROW = 3
FIRST_DATA_ROW = 4
PREFILLED_SHEETS = ("01_Run_Summary", "02_Run_Config")
RUN_SHEETS = (
    "01_Run_Summary", "02_Run_Config", "03_Clustering_History", "04_Supervised_History",
    "05_Correlation_Matrix", "06_Quantum_Labels", "07_Pairwise_Distance", "08_Dimension_Activity",
    "09_Per_Class_Metrics", "10_Confusion_Data", "11_Test_Predictions", "19_Threshold_Sensitivity",
)
ID_COLUMNS = {
    C.TRACK_PRIMARY: ("condition_id", "run_uid"),
    C.TRACK_ABLATION: ("ablation_condition_id", "ablation_run_uid"),
    C.TRACK_MORE_REFERENCE: ("reference_condition_id", "reference_run_uid"),
}


def _fmt(v):
    if isinstance(v, (list, tuple, dict)):
        import json

        return json.dumps(v, ensure_ascii=False)
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating,)):
        return float(v)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    return v


class RunArtifacts:
    """Lazy reader of one run folder."""

    def __init__(self, run_dir: str | Path):
        self.dir = Path(run_dir)

    def _p(self, rel: str) -> Path:
        return self.dir / rel

    @cached_property
    def manifest(self) -> dict:
        return load_json(self._p("config.json"))

    @cached_property
    def status(self) -> dict | None:
        p = self._p("logs/run_status.json")
        return load_json(p) if p.exists() else None

    @cached_property
    def started(self) -> dict:
        return load_json(self._p("logs/run_started.json"))

    @cached_property
    def stage_wall(self) -> dict:
        return {r["stage"]: r["wall_sec"] for r in read_jsonl(self._p("logs/stage_timing.jsonl"))}

    @cached_property
    def clustering_log(self) -> list:
        return read_jsonl(self._p("logs/clustering_log.jsonl"))

    @cached_property
    def supervised_log(self) -> list:
        return read_jsonl(self._p("logs/supervised_log.jsonl"))

    @cached_property
    def opt_c(self) -> dict:
        return load_json(self._p("logs/clustering_optimizer_result.json"))

    @cached_property
    def opt_s(self) -> dict:
        return load_json(self._p("logs/supervised_optimizer_result.json"))

    @cached_property
    def metrics(self) -> dict:
        return load_json(self._p("logs/metrics_final.json"))

    @cached_property
    def confusion(self) -> np.ndarray:
        return load_npy(self._p("logs/confusion_matrix.npy"))

    @cached_property
    def quantum_labels(self) -> dict:
        return load_json(self._p("artifacts/quantum_labels.json"))

    @cached_property
    def label_diag(self) -> dict:
        return load_json(self._p("artifacts/quantum_label_diagnostics.json"))

    @cached_property
    def structural(self) -> dict:
        return load_json(self._p("artifacts/structural_diagnostics.json"))

    @cached_property
    def s_matrix(self) -> np.ndarray:
        return load_npy(self._p("artifacts/correlation_matrix.npy"))

    @cached_property
    def raw_mse(self) -> np.ndarray:
        return load_npy(self._p("artifacts/correlation_raw_mse.npy"))

    def npy(self, rel: str) -> np.ndarray:
        return load_npy(self._p(rel))


def _common(run: RunArtifacts) -> dict:
    m = run.manifest
    return {
        "condition_id": m["condition_id"], "run_uid": m["run_uid"], "seed": m["seed"],
        "K": m["n_classes"], "feature_method": m["feature_method"], "architecture": m["architecture"],
    }


def _row_at(log: list, eval_id: int) -> dict:
    row = log[eval_id]
    if row["eval_id"] != eval_id:
        raise AssertionError("log rows must be indexed by eval_id")
    return row


def build_run_summary(run: RunArtifacts) -> dict:
    m, oc, os_, mf = run.manifest, run.opt_c, run.opt_s, run.metrics
    crow = _row_at(run.clustering_log, oc["final_point_eval_id"])
    srow = _row_at(run.supervised_log, os_["final_point_eval_id"])
    cm = run.confusion
    sw = run.stage_wall
    t = m["timing"]
    sd = run.structural["checkpoints"]
    first = "clustering_final" if "clustering_final" in sd else "clustering_selected"
    k = m["n_classes"]
    row = {
        "condition_id": m["condition_id"], "run_uid": m["run_uid"], "run_name": m["run_name"],
        "status": run.status["status"] if run.status else None, "run_mode": m["run_mode"], "K": k,
        "classes": ",".join(str(c) for c in m["classes"]), "feature_method": m["feature_method"],
        "architecture": m["architecture"], "seed": m["seed"], "replication_id": f"S{m['seed']}",
        "n_params": oc["n_params"], "n_train_total": m["n_train_per_class"] * k,
        "n_val_total": m["n_val_per_class"] * k, "n_test_total": m["n_test_per_class"] * k,
        "split_manifest": m["split_manifest"], "attempt": m["attempt"],
        "cluster_nfev": oc["nfev"], "final_cluster_loss": oc["fun"],
        "cluster_final_point_eval_id": oc["final_point_eval_id"],
        "cluster_best_observed_eval_id": oc["best_observed_eval_id"],
        "cluster_n_optimization_evals": oc["n_optimization_evals"], "cluster_optimizer_success": oc["success"],
        "cluster_optimizer_status": oc["status"], "cluster_optimizer_message": oc["message"],
        "pseudo_accuracy_val": crow["pseudo_accuracy_val"], "avg_margin_val": crow["avg_margin_val"],
        "min_separation": crow["min_separation"], "min_separation_ratio": crow["min_separation_ratio"],
        "correlation_consistency": crow["correlation_consistency"], "active_dimensions": crow["active_dimensions"],
        "yodd_norm_fraction": sd[first]["yodd_norm_fraction"],
        "supervised_nfev": os_["nfev"], "final_supervised_train_loss": os_["fun"],
        "final_supervised_val_loss": srow["val_loss"], "supervised_final_point_eval_id": os_["final_point_eval_id"],
        "supervised_best_observed_eval_id": os_["best_observed_eval_id"],
        "supervised_n_optimization_evals": os_["n_optimization_evals"],
        "supervised_optimizer_success": os_["success"], "supervised_optimizer_status": os_["status"],
        "supervised_optimizer_message": os_["message"],
        "accuracy": mf["accuracy"], "precision_macro": mf["precision"], "recall_macro": mf["recall"],
        "f1_macro": mf["f1_score"], "n_correct": int(np.trace(cm)), "n_wrong": int(cm.sum() - np.trace(cm)),
        "total_runtime_sec": m["total_runtime_sec"], "degenerate_quantum_label": m["degenerate_quantum_label"],
        "n_degenerate_test_outputs": mf["n_degenerate_test_outputs"],
        "started_at": run.started["started_at"], "finished_at": t["finished_at"], "total_cpu_sec": t["total_cpu_sec"],
        "data_pipeline_sec": sw.get("data_pipeline"), "setup_sec": sw.get("setup"),
        "clustering_loop_sec": sw.get("clustering_loop"),
        "quantum_label_extraction_sec": sw.get("quantum_label_extraction"),
        "supervised_loop_sec": sw.get("supervised_loop"), "final_evaluation_sec": sw.get("final_evaluation"),
        "structural_diagnostics_sec": sw.get("structural_diagnostics"),
        "n_parallel_declared": m["n_parallel_declared"], "load_avg_1m_start": run.started["load_avg_1m_start"],
        "load_avg_1m_end": t["load_avg_1m_end"], "hostname": run.started["hostname"],
        "cpu_count": run.started["cpu_count"], "thread_env": run.started["thread_env"],
        "yodd_norm_fraction_supervised": sd["supervised_final"]["yodd_norm_fraction"],
        "yodd_max_abs_clustering": sd[first]["yodd_max_abs"],
        "min_separation_val_ratio": crow["min_separation_val_ratio"],
        "git_commit": m["environment"]["git_commit"],
        "backend_max_abs_diff": m["backend_self_check"]["max_abs_diff"],
    }
    if m["track"] != C.TRACK_PRIMARY:
        row.update({
            "model_code": m["model_code"], "matching_run_uid_A": m.get("matching_run_uid_A"),
            "matching_run_uid_D": m.get("matching_run_uid_D"), "n_trainable_params": m["n_trainable_params"],
            "n_fixed_rz_params": m["n_fixed_rz_params"], "n_variational_layers": m["n_variational_layers"],
            "variational_gate_count": m["variational_gate_count"],
            "entangling_block_count": m["entangling_block_count"], "cnot_count": m["cnot_count"],
            "circuit_depth": m["circuit_resources"]["depth"], "ry_core_seed": m["ry_core_seed"],
            "ry_extra_seed": m["ry_extra_seed"] if m["architecture"] == C.ARCH_MORE_HD_60P else None,
            "rz_phase_seed": m["rz_phase_seed"] if m["architecture"] == C.ARCH_MORE_HD_C_FIXED_RZ else None,
            "more_init_seed": m["more_init_seed"] if m["architecture"] == C.ARCH_MORE_REPRO else None,
            "fixed_rz_sha256": load_json(run._p("artifacts/initialization_manifest.json")).get("fixed_rz_sha256"),
            "paired_pair_manifest_sha256": m["paired_pair_manifest_sha256"],
            "paired_split_manifest_sha256": m["split_manifest_sha256"],
            "paired_input_check": m["paired_input_check"]["status"],
        })
    return row


def build_run_config(run: RunArtifacts) -> dict:
    m = run.manifest
    return {
        "condition_id": m["condition_id"], "run_uid": m["run_uid"], "run_name": m["run_name"],
        "run_mode": m["run_mode"], "K": m["n_classes"], "classes": ",".join(str(c) for c in m["classes"]),
        "n_train_per_class": m["n_train_per_class"], "n_val_per_class": m["n_val_per_class"],
        "n_test_per_class": m["n_test_per_class"], "architecture": m["architecture"],
        "feature_method": m["feature_method"], "n_data_qubits": m["n_data_qubits"],
        "n_readout_qubits": m["n_readout_qubits"], "n_input_channels": m["n_input_channels"],
        "pca_n_components": m["pca_n_components"], "hu_raw_dim": m["hu_raw_dim"],
        "hu_padding_value": m["hu_padding_value"], "zernike_terms": m["zernike_terms"],
        "cobyla_tol": m["cobyla_tol"], "cobyla_rhobeg": m["cobyla_rhobeg"],
        "max_nfev_clustering": m["max_nfev_clustering"], "max_nfev_supervised": m["max_nfev_supervised"],
        "scipy_version": m["scipy_version"], "master_seed": m["master_seed"], "data_seed": m["data_seed"],
        "ry_core_seed": m["ry_core_seed"], "rz_phase_seed": m["rz_phase_seed"], "pair_seed": m["pair_seed"],
        "label_seed": m["label_seed"], "split_manifest": m["split_manifest"],
        "paired_split_key": m["paired_split_key"], "nested_class_split": True,
        "active_dim_threshold": m["active_dim_threshold"], "n_params": run.opt_c["n_params"],
        "subseed_rule": m["subseed_rule"], "numpy_version": m["numpy_version"],
        "pennylane_version": m["pennylane_version"], "simulation_mode": m["simulation_mode"],
        "shots": m["shots"], "sim_dtype": m["sim_dtype"], "device_name": m["device_name"],
        "eps_norm": m["eps_norm"], "centroid_rule": m["centroid_rule"], "centroid_source": m["centroid_source"],
        "val_monitor_policy": m["val_monitor_policy"], "n_val_monitor_per_class": m["n_val_monitor_per_class"],
        "correlation_rule": m["correlation_rule"], "pca_svd_solver": m["pca_svd_solver"],
    }


def build_history(run: RunArtifacts, which: str) -> list:
    base = _common(run)
    log = run.clustering_log if which == "clustering" else run.supervised_log
    return [{**base, **row} for row in log]


def build_correlation(run: RunArtifacts) -> list:
    base = _common(run)
    classes = run.manifest["classes"]
    s, mse = run.s_matrix, run.raw_mse
    return [{**base, "class_i": ci, "class_j": cj, "raw_mse": float(mse[i, j]), "S_value": float(s[i, j])}
            for i, ci in enumerate(classes) for j, cj in enumerate(classes)]


def _labels_array(run: RunArtifacts) -> tuple[np.ndarray, list]:
    ql = run.quantum_labels
    obs = ql["observables"]
    classes = run.manifest["classes"]
    return np.array([[ql["labels"][str(c)][o] for o in obs] for c in classes]), obs


def build_quantum_labels(run: RunArtifacts) -> list:
    base = _common(run)
    labels, obs = _labels_array(run)
    rows = []
    for i, c in enumerate(run.manifest["classes"]):
        d = run.label_diag["per_class"][str(c)]
        row = {**base, "class": c}
        row.update({o: float(labels[i, j]) for j, o in enumerate(obs)})
        row.update({"n_label_samples": d["n_samples"], "n_degenerate_samples_excluded": d["n_degenerate_samples_excluded"],
                    "label_prenorm": d["prenorm"], "label_degenerate": d["degenerate"]})
        rows.append(row)
    return rows


def build_pairwise(run: RunArtifacts) -> list:
    base = _common(run)
    labels, _ = _labels_array(run)
    classes = run.manifest["classes"]
    k = len(classes)
    dmat = pairwise_label_distances(labels, run.manifest["eps_norm"])
    bound = C.simplex_bound(k)
    return [{**base, "class_i": classes[i], "class_j": classes[j], "S_value": float(run.s_matrix[i, j]),
             "cosine_distance": float(dmat[i, j]), "simplex_bound": bound, "separation_ratio": float(dmat[i, j] / bound)}
            for i in range(k) for j in range(i + 1, k)]


def build_dimension_activity(run: RunArtifacts) -> list:
    base = _common(run)
    thr = run.manifest["active_dim_threshold"]
    rows = []
    for name, ck in run.structural["checkpoints"].items():
        for p in ck["per_observable"]:
            rows.append({**base, "checkpoint": name, "observable": p["observable"], "is_yodd": p["is_yodd"],
                         "mean_abs": p["mean_abs"], "max_abs": p["max_abs"], "std_abs": p["std_abs"],
                         "active_threshold": thr, "is_active": p["is_active"], "norm_fraction": p["norm_fraction"]})
    return rows


def build_per_class(run: RunArtifacts) -> list:
    base = _common(run)
    return [{**base, **pc} for pc in run.metrics["per_class"]]


def build_confusion(run: RunArtifacts) -> list:
    base = _common(run)
    classes = run.manifest["classes"]
    cm = run.confusion
    return [{**base, "actual_class": a, "predicted_class": b, "count": int(cm[i, j])}
            for i, a in enumerate(classes) for j, b in enumerate(classes)]


def build_test_predictions(run: RunArtifacts) -> list:
    base = _common(run)
    y = run.npy("artifacts/y_test.npy")
    idx = run.npy("logs/test_mnist_index.npy")
    pred = run.npy("logs/test_predictions.npy")
    dist = run.npy("logs/test_label_distances.npy")
    margin = run.npy("logs/test_margins.npy")
    classes = run.manifest["classes"]
    col = {c: i for i, c in enumerate(classes)}
    rows = []
    for s in range(y.size):
        rows.append({**base, "sample_id": s, "mnist_index": int(idx[s]), "true_class": int(y[s]),
                     "pred_class": int(pred[s]), "distance_true": float(dist[s, col[int(y[s])]]),
                     "distance_pred": float(dist[s, col[int(pred[s])]]), "margin": float(margin[s]),
                     "correct": bool(pred[s] == y[s])})
    return rows


def build_threshold_sensitivity(run: RunArtifacts) -> list:
    base = _common(run)
    thr = run.manifest["active_dim_threshold"]
    rows = []
    for name, ck in run.structural["checkpoints"].items():
        for t in ck["threshold_sensitivity"]:
            rows.append({**base, "checkpoint": name, "threshold": t["threshold"],
                         "is_primary_threshold": t["threshold"] == thr, "n_active": t["n_active"],
                         "n_active_yodd": t["n_active_yodd"]})
    return rows


def rows_for_run(run_dir) -> dict:
    run = RunArtifacts(run_dir)
    return {
        "01_Run_Summary": [build_run_summary(run)],
        "02_Run_Config": [build_run_config(run)],
        "03_Clustering_History": build_history(run, "clustering"),
        "04_Supervised_History": build_history(run, "supervised"),
        "05_Correlation_Matrix": build_correlation(run),
        "06_Quantum_Labels": build_quantum_labels(run),
        "07_Pairwise_Distance": build_pairwise(run),
        "08_Dimension_Activity": build_dimension_activity(run),
        "09_Per_Class_Metrics": build_per_class(run),
        "10_Confusion_Data": build_confusion(run),
        "11_Test_Predictions": build_test_predictions(run),
        "19_Threshold_Sensitivity": build_threshold_sensitivity(run),
    }


# ---------------------------------------------------------------------------
# workbook helpers
# ---------------------------------------------------------------------------


def headers(ws) -> list:
    return [c.value for c in ws[HEADER_ROW] if c.value is not None]


def _rename_ids(row: dict, track: str) -> dict:
    cid, ruid = ID_COLUMNS[track]
    if cid == "condition_id":
        return row
    out = dict(row)
    if "condition_id" in out:
        out[cid] = out.pop("condition_id")
    if "run_uid" in out:
        out[ruid] = out.pop("run_uid")
    return out


def _row_values(hdr: list, row: dict) -> list:
    return [_fmt(row.get(h)) for h in hdr]


def validate_schema_map(wb, allow_open: bool = False) -> list:
    """VALIDATE_SCHEMA_MAP: (1) every column mapped, (2) no stale map rows, (3) no GAP/PENDING."""
    problems = []
    sm = wb["00_Schema_Map"]
    mapped = {}
    for r in range(FIRST_DATA_ROW, sm.max_row + 1):
        sheet, col, _, _, _, status, _ = [c.value for c in sm[r]][:7]
        if sheet is None:
            continue
        mapped[(sheet, col)] = status
    for ws in wb.worksheets:
        if ws.title in ("00_README", "00_Schema_Map") or ws.title.startswith("99_"):
            continue
        for h in headers(ws):
            if (ws.title, h) not in mapped:
                problems.append(f"unmapped column {ws.title}.{h}")
    for (sheet, col), status in mapped.items():
        if sheet not in wb.sheetnames or col not in headers(wb[sheet]):
            problems.append(f"stale schema row {sheet}.{col}")
        if not allow_open and status in ("GAP", "PENDING"):
            problems.append(f"open schema row {sheet}.{col} ({status})")
    return problems


def update_local_run_spreadsheet(run_dir) -> Path:
    """UPDATE_LOCAL_RUN_SPREADSHEET: fill only this run's rows in runs/<run>/run_result.xlsx."""
    run_dir = Path(run_dir)
    path = run_dir / C.LOCAL_RUN_SPREADSHEET_NAME
    if not path.exists():
        raise FileNotFoundError(path)
    track = load_json(run_dir / "config.json")["track"]
    wb = openpyxl.load_workbook(path)
    problems = validate_schema_map(wb)
    if problems:
        raise ValueError("schema map invalid: " + "; ".join(problems[:10]))
    for h in (ws for ws in wb.worksheets):
        if "Iteration" in headers(h):
            raise ValueError("legacy header 'Iteration' (G1-02)")
    rows = rows_for_run(run_dir)
    for sheet, items in rows.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        hdr = headers(ws)
        id_col = ID_COLUMNS[track][1]
        for item in items:
            item = _rename_ids(item, track)
            values = _row_values(hdr, item)
            target = None
            if sheet in PREFILLED_SHEETS and id_col in hdr:
                ci = hdr.index(id_col) + 1
                for r in range(FIRST_DATA_ROW, ws.max_row + 1):
                    if ws.cell(r, ci).value == item[id_col]:
                        target = r
                        break
            if target is None:
                ws.append(values)
            else:
                for j, v in enumerate(values, start=1):
                    if v is not None:
                        ws.cell(target, j, v)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# consolidator (G3-03)
# ---------------------------------------------------------------------------


def discover_runs(runs_root, track: str, run_mode: str | None = None) -> list[Path]:
    out = []
    for cfg_path in sorted(Path(runs_root).rglob("config.json")):
        try:
            m = load_json(cfg_path)
        except Exception:
            continue
        if m.get("track") != track or "run_uid" not in m:
            continue
        if run_mode is not None and m.get("run_mode") != run_mode:
            continue
        out.append(cfg_path.parent)
    return out


def select_completed_attempts(run_dirs: list[Path]) -> dict:
    """Exactly one COMPLETED attempt per run_uid; duplicates are an error (key = run_uid + attempt)."""
    by_uid: dict = {}
    seen_keys = set()
    for d in run_dirs:
        m = load_json(d / "config.json")
        st = load_json(d / "logs" / "run_status.json")
        key = (m["run_uid"], m["attempt"])
        if key in seen_keys:
            raise ValueError(f"duplicate run_uid+attempt {key}")
        seen_keys.add(key)
        if st["status"] != C.STATUS_COMPLETED:
            continue
        if m["run_uid"] in by_uid:
            raise ValueError(f"run_uid {m['run_uid']} has more than one COMPLETED attempt")
        by_uid[m["run_uid"]] = d
    return by_uid


def benchmark_rows(benchmark_files) -> list:
    rows = []
    for f in benchmark_files:
        d = load_json(f)
        meta = {k: d.get(k) for k in ("device_name", "sim_dtype", "pennylane_version", "numpy_version",
                                      "hostname", "cpu_model", "thread_env", "timestamp")}
        for r in d["rows"]:
            rows.append({"benchmark_file": Path(f).name, **r, **meta})
    return rows


def baseline_rows(baseline_csv, primary_rows: dict) -> list:
    """Rows of sheet 15 from runs/baselines*/baseline_results.csv, joined with A/D summaries."""
    import csv

    out = []
    with open(baseline_csv, encoding="utf-8") as fh:
        for b in csv.DictReader(fh):
            k = int(b["K"])
            a = primary_rows.get(b["source_run_uid_A"])
            d = primary_rows.get(b["check_run_uid_D"]) if b["check_run_uid_D"] else None
            acc = float(b["accuracy"])
            f1 = float(b["macro_f1"])
            row = {"K": k, "classes": ",".join(str(c) for c in range(k)), "feature_method": b["feature_method"],
                   "seed": int(b["seed"]), "baseline": b["baseline"], "source_run_uid_A": b["source_run_uid_A"],
                   "check_run_uid_D": b["check_run_uid_D"] or None,
                   "input_hash_match_A_D": None if b["input_hash_match_A_D"] in ("", "None") else b["input_hash_match_A_D"] == "True",
                   "C_selected": b["C_selected"] or "n/a", "lr_converged": b["lr_converged"] or "n/a",
                   "accuracy": acc, "macro_f1": f1, "chance_1_over_K": 1.0 / k, "status": "COMPLETED"}
            if a:
                row.update({"accuracy_more_hd": a["accuracy"], "f1_more_hd": a["f1_macro"],
                            "delta_acc_more_hd_minus_baseline": a["accuracy"] - acc,
                            "delta_f1_more_hd_minus_baseline": a["f1_macro"] - f1})
            if d:
                row.update({"accuracy_more_hd_c": d["accuracy"], "f1_more_hd_c": d["f1_macro"],
                            "delta_acc_more_hd_c_minus_baseline": d["accuracy"] - acc,
                            "delta_f1_more_hd_c_minus_baseline": d["f1_macro"] - f1})
            out.append(row)
    return out


def consolidate(template_path, run_dirs: list[Path], out_path, track: str, jalur_b_dirs=(), benchmark_files=(),
                baseline_csv=None) -> dict:
    """Rebuild a consolidated workbook from the template and run artifacts (idempotent).

    The output is always regenerated from scratch, so repeating the consolidation
    with the same runs yields the same content and no run can alter another
    run's rows. Large append-only sheets are streamed (openpyxl write-only).
    """
    selected = select_completed_attempts(run_dirs)
    tpl = openpyxl.load_workbook(template_path)
    problems = validate_schema_map(tpl)
    if problems:
        raise ValueError("template schema map invalid: " + "; ".join(problems[:10]))
    per_sheet: dict = {name: [] for name in RUN_SHEETS}
    log_rows = []
    summaries = {}
    for uid in sorted(selected):
        d = selected[uid]
        rows = rows_for_run(d)
        summaries[uid] = rows["01_Run_Summary"][0]
        for sheet, items in rows.items():
            per_sheet[sheet].extend(_rename_ids(it, track) for it in items)
        m = load_json(d / "config.json")
        log_rows.append([uid, m["attempt"], str(d), sha256_file(d / "config.json")])
    if track == C.TRACK_PRIMARY:
        from .jalur_b import jalur_b_row

        per_sheet["17_JalurB_Selection"] = [jalur_b_row(d) for d in sorted(jalur_b_dirs)]
        per_sheet["18_Circuit_Benchmark"] = benchmark_rows(sorted(benchmark_files))
        if baseline_csv is not None:
            per_sheet["15_Classical_Baselines"] = baseline_rows(baseline_csv, summaries)
    out = openpyxl.Workbook(write_only=True)
    id_col = ID_COLUMNS[track][1]
    for ws_t in tpl.worksheets:
        ws = out.create_sheet(ws_t.title)
        tpl_rows = list(ws_t.iter_rows(values_only=True))
        for r in tpl_rows[:HEADER_ROW]:
            ws.append(list(r))
        hdr = headers(ws_t)
        body = tpl_rows[HEADER_ROW:]
        if ws_t.title == "15_Classical_Baselines" and per_sheet.get(ws_t.title):
            key_cols = ("K", "feature_method", "seed", "baseline")
            fill = {tuple(row[c] for c in key_cols): row for row in per_sheet[ws_t.title]}
            used = set()
            for r in body:
                r = list(r)[: len(hdr)]
                key = tuple(r[hdr.index(c)] for c in key_cols)
                if key in fill:
                    r = [v if v is not None else r[i] for i, v in enumerate(_row_values(hdr, fill[key]))]
                    used.add(key)
                ws.append(r)
            for key, row in fill.items():
                if key not in used:
                    ws.append(_row_values(hdr, row))
        elif ws_t.title in PREFILLED_SHEETS:
            fill = {row[id_col]: row for row in per_sheet[ws_t.title]}
            used = set()
            for r in body:
                r = list(r)[: len(hdr)]
                key = r[hdr.index(id_col)] if id_col in hdr else None
                if key in fill:
                    r = [v if v is not None else r[i] for i, v in enumerate(_row_values(hdr, fill[key]))]
                    used.add(key)
                ws.append(r)
            for key, row in fill.items():           # e.g. pilot runs absent from the template
                if key not in used:
                    ws.append(_row_values(hdr, row))
        else:
            for r in body:
                ws.append(list(r))
            for row in per_sheet.get(ws_t.title, []):
                ws.append(_row_values(hdr, row))
    log = out.create_sheet("99_Consolidation_Log")
    log.append(["Consolidation log — one row per consolidated run (operational, not analysed)"])
    log.append([])
    log.append(["run_uid", "attempt", "run_dir", "config_json_sha256"])
    for r in log_rows:
        log.append(r)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = out_path.with_suffix(".tmp.xlsx")
    out.save(tmp)
    shutil.move(tmp, out_path)
    return {"n_runs": len(selected), "run_uids": sorted(selected), "out_path": str(out_path)}
