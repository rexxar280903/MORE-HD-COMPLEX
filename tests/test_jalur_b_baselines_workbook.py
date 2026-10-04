"""G0-02 (Jalur B selector), G1-10 (classical baselines), G3-01/02/03/06/07 (workbooks)."""

import shutil

import numpy as np
import openpyxl
import pytest

from core import constants as C
from core.artifact_io import load_json, read_jsonl
from core.jalur_b import select_clustering_checkpoint
from tests.conftest import ROOT, make_project, requires_mnist, tiny_config


def _row(eid, phase, acc, msv, mtd, loss, deg=0):
    return {"eval_id": eid, "phase": phase, "pseudo_accuracy_val": acc, "min_separation_val": msv,
            "mean_true_distance_val": mtd, "train_loss": loss, "n_degenerate_centroids_train": deg,
            "n_degenerate_centroids_val": 0, "active_dimensions": 15 - eid % 3, "avg_margin_val": eid}


def test_selector_rules_and_determinism():
    n = 2
    rows = [_row(0, "initial_simplex", 0.99, 1, 0, 0), _row(1, "initial_simplex", 0.99, 1, 0, 0),
            _row(2, "initial_simplex", 0.99, 1, 0, 0),
            _row(3, "optimization", 0.80, 0.5, 0.2, -1.0), _row(4, "optimization", 0.90, 0.4, 0.3, -0.5),
            _row(5, "optimization", 0.90, 0.6, 0.3, -0.4), _row(6, "optimization", 0.90, 0.6, 0.2, -0.3),
            _row(7, "optimization", 0.90, 0.6, 0.2, -0.3), _row(8, "optimization", 0.95, 0.9, 0.1, -2.0, deg=1)]
    sel, rec, cands = select_clustering_checkpoint(rows, n)
    assert sel == 6                                     # acc -> sep -> distance -> loss -> earliest eval_id
    assert rec["eligible_eval_ids"] == [3, 4, 5, 6, 7]  # simplex and degenerate evaluations excluded
    assert rec["test_used_for_selection"] is False
    sel2, rec2, cands2 = select_clustering_checkpoint(list(reversed(rows)), n)
    assert sel2 == sel and [c["eval_id"] for c in cands2] == [c["eval_id"] for c in cands]
    with pytest.raises(ValueError):
        select_clustering_checkpoint(rows[:3], n)
    with pytest.raises(AssertionError):
        select_clustering_checkpoint([_row(0, "optimization", 0.5, 0.5, 0.5, 0.0)], n)


@requires_mnist
def test_jalur_b_end_to_end(tmp_path, mnist, monkeypatch):
    from core.evaluation import OfficialTestVault
    from core.jalur_b import jalur_b_row, run_jalur_b
    from core.pipeline import run_condition

    proj = make_project(tmp_path)
    cfg = tiny_config(proj, "MORE-HD", "PCA", 3)
    run_condition(cfg, mnist=mnist, write_workbook=False)
    opened_in = []
    orig = OfficialTestVault.open

    def spy(self):
        opened_in.append(self.clock.current_stage)
        return orig(self)

    monkeypatch.setattr(OfficialTestVault, "open", spy)
    m = run_jalur_b(cfg.run_dir, n_parallel_declared=1, project_root=proj)
    assert opened_in == ["final_evaluation"]
    sel = load_json(cfg.run_dir / "logs" / "selected_checkpoint.json")
    assert sel["selected_eval_id"] == m["selected_eval_id"] and sel["selected_metrics"]["phase"] == "optimization"
    assert m["max_nfev_supervised"] == cfg.max_nfev_supervised
    jb_dir = proj / "runs" / m["run_name"]
    stages = [s["stage"] for s in read_jsonl(jb_dir / "logs" / "stage_timing.jsonl")]
    assert stages == list(C.STAGES_JALUR_B)
    row = jalur_b_row(jb_dir)
    assert row["test_used_for_selection"] is False and row["selected_phase"] == "optimization"
    from core.jalur_b import select_from_source

    assert select_from_source(cfg.run_dir)[0] == m["selected_eval_id"]          # reproducible


@requires_mnist
def test_classical_baseline_cell(tmp_path, mnist):
    from core.classical_baselines import run_baseline_cell
    from core.pipeline import run_condition

    proj = make_project(tmp_path)
    for arch in ("MORE-HD", "MORE-HD-C"):
        run_condition(tiny_config(proj, arch, "PCA", 3), mnist=mnist, write_workbook=False)
    log = []
    rows = run_baseline_cell(proj, 42, "PCA", 3, proj / "runs" / "baselines_pilot", (20, 6, 8), access_log=log)
    frozen = log.index("MODELS_FROZEN")
    assert "X_test_scaled" in log[frozen:] and "X_test_scaled" not in log[:frozen]
    names = [r["baseline"] for r in rows]
    assert names == ["NC", "LR", "CHANCE"]
    lr = rows[1]
    assert lr["C_selected"] in (0.01, 0.1, 1.0, 10.0, 100.0) and lr["input_hash_match_A_D"] is True
    cv = load_json(proj / "runs" / "baselines_pilot" / "S42_cls-0-1-2_PCA" / "LR_cv_log.json")
    best = max(g["val_accuracy"] for g in cv["grid"])
    assert cv["C_selected"] == min(g["C"] for g in cv["grid"] if g["val_accuracy"] == best)
    assert rows[2]["accuracy"] == pytest.approx(1 / 3)


def test_templates_have_complete_schema_maps():
    from core.workbook import validate_schema_map

    for path in (C.MASTER_SPREADSHEET_PATH, C.ABLATION_SPREADSHEET_PATH, C.MORE_REFERENCE_SPREADSHEET_PATH):
        wb = openpyxl.load_workbook(ROOT / path)
        assert validate_schema_map(wb) == [], path
        for ws in wb.worksheets:
            assert "Iteration" not in [c.value for c in ws[3]]


ALLOWED_NULL = {
    "n_parallel_declared", "load_avg_1m_start", "load_avg_1m_end", "git_commit", "correlation_consistency",
    "pca_n_components", "hu_raw_dim", "hu_padding_value", "zernike_terms", "pca_svd_solver",
    "n_val_monitor_per_class", "shots", "ry_extra_seed", "rz_phase_seed", "more_init_seed", "fixed_rz_sha256",
}


@requires_mnist
def test_per_run_workbook_and_consolidator(tmp_path, mnist):
    from core.pipeline import run_condition
    from core.workbook import consolidate, discover_runs, validate_schema_map

    proj = make_project(tmp_path)
    cfg = tiny_config(proj, "MORE-HD-C", "ZERNIKE", 3, nfev=320)        # long history (G3-01)
    run_condition(cfg, mnist=mnist, write_workbook=True)
    run_condition(tiny_config(proj, "MORE-HD", "ZERNIKE", 3), mnist=mnist, write_workbook=True)
    wb = openpyxl.load_workbook(cfg.run_dir / "run_result.xlsx", read_only=True)
    hist = list(wb["03_Clustering_History"].iter_rows(values_only=True))
    assert len(hist) - 3 == 320                                        # every evaluation, no truncation
    assert len(list(wb["04_Supervised_History"].iter_rows(values_only=True))) - 3 == 320
    summary = list(wb["01_Run_Summary"].iter_rows(values_only=True))
    hdr = summary[2]
    mine = [r for r in summary[3:] if r[1] == cfg.run_uid]
    assert len(mine) == 1
    missing = [h for h, v in zip(hdr, mine[0]) if v is None and h not in ALLOWED_NULL]
    assert missing == []
    preds = list(wb["11_Test_Predictions"].iter_rows(values_only=True))
    assert len(preds) - 3 == 3 * cfg.n_test_per_class and None not in preds[3]
    wb.close()
    # consolidation: idempotent, duplicate-safe, primary only
    dirs = discover_runs(proj / "runs", C.TRACK_PRIMARY)
    out1 = consolidate(ROOT / C.MASTER_SPREADSHEET_PATH, dirs, tmp_path / "c1.xlsx", C.TRACK_PRIMARY)
    out2 = consolidate(ROOT / C.MASTER_SPREADSHEET_PATH, dirs, tmp_path / "c2.xlsx", C.TRACK_PRIMARY)
    assert out1["n_runs"] == 2
    w1 = openpyxl.load_workbook(tmp_path / "c1.xlsx", read_only=True)
    w2 = openpyxl.load_workbook(tmp_path / "c2.xlsx", read_only=True)
    for name in w1.sheetnames:
        assert list(w1[name].iter_rows(values_only=True)) == list(w2[name].iter_rows(values_only=True)), name
    assert validate_schema_map(openpyxl.load_workbook(tmp_path / "c1.xlsx")) == []
    dup = proj / "runs" / (cfg.run_dir.name + "_copy")
    shutil.copytree(cfg.run_dir, dup)
    with pytest.raises(ValueError):
        consolidate(ROOT / C.MASTER_SPREADSHEET_PATH, discover_runs(proj / "runs", C.TRACK_PRIMARY),
                    tmp_path / "c3.xlsx", C.TRACK_PRIMARY)
